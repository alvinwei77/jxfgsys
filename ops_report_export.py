# -*- coding: utf-8 -*-
"""经营报告（分期与专题）—— 一键导出 Excel。

  GET /api/export/ops_report                全部期（缺省），统计报表页「导出经营报告 Excel」按钮走这里
  GET /api/export/ops_report?periods=2026-10  只出指定期
  命令行：python3 ops_report_export.py [--periods 2026-08,2026-10] [--out 文件名.xlsx]

内容：①总览 / ②各期明细(每期一个 sheet) / 调货人工资 / 量补汇总 / 票源额度水位(A1配票·A2反向·B自带票，
额度取「票源额度分期」生效值)。数据源 = 系统当前库（只读），逐车走 engine 实时口径，与页面一致。
（2026-10-07 由工作区脚本 gen_ops_report_any.py 移植进系统，逻辑不变。）
"""
import argparse
import datetime
import json
import sqlite3
import sys

import os  # noqa: E402

import engine  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # noqa: E402
from openpyxl.formatting.rule import DataBarRule  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))


def _db_path():
    """与系统同一套取库路径规则（优先 SCRAPMETAL_DB，其次 SCRAPMETAL_TEST_DB，最后仓库自带库）。

    2026-10-09 修：之前只认 SCRAPMETAL_DB，于是拿 SCRAPMETAL_TEST_DB 跑副本/回归时，
    这个导出会静默去读仓库自带的 data/scrapmetal.db（副本里没有该文件就直接 500）。
    """
    return (os.environ.get("SCRAPMETAL_DB") or os.environ.get("SCRAPMETAL_TEST_DB")
            or os.path.join(HERE, "data", "scrapmetal.db"))


REPORTS_DIR = os.path.join(HERE, "data", "reports")


def _high_water_pct():
    """「水位偏高」判定线，返回小数（0.8 = 80%）。取自「系统 → 参数设置」的 quota_high_water_pct，
    读不到就回退原来写死的 0.8（C档 2026-10-09）。"""
    try:
        conn = sqlite3.connect(f"file:{_db_path()}?mode=ro", uri=True)
        try:
            import settings
            return float(settings.get(conn, "quota_high_water_pct", 80.0)) / 100.0
        finally:
            conn.close()
    except Exception:
        return 0.8

TITLE = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
SUB   = Font(name="微软雅黑", size=9, color="595959")
SB    = Font(name="微软雅黑", size=11, bold=True)
H     = Font(name="微软雅黑", size=10, bold=True)
N     = Font(name="微软雅黑", size=10)
SM    = Font(name="微软雅黑", size=9)
NT    = Font(name="微软雅黑", size=10)
MT    = Font(name="微软雅黑", size=10)
M0    = Font(name="微软雅黑", size=10)

FILLT = PatternFill("solid", fgColor="1F4E79")
FILLH = PatternFill("solid", fgColor="D9E2F3")
GRY   = PatternFill("solid", fgColor="F2F2F2")
GRN   = PatternFill("solid", fgColor="E2EFDA")
BLU   = PatternFill("solid", fgColor="DDEBF7")
YEL   = PatternFill("solid", fgColor="FFF2CC")
BD    = Border(*[Side(style="thin", color="BFBFBF")] * 4)
CT = Alignment(horizontal="center", vertical="center", wrap_text=True)
LT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RT = Alignment(horizontal="right", vertical="center")

CHAIN_ORDER = ["配票", "反向", "直签带票"]
AGG_KEYS = ("B", "D", "give", "front_inv", "front_paid", "pa", "rw", "cm", "sfee", "net", "ires", "adv")


def _put_one(ws, r, c, v, font, fill, align, fmt, bold):
    cell = ws.cell(r, c, v)
    cell.font = Font(name="微软雅黑", size=10, bold=True) if bold else font
    if fill:
        cell.fill = fill
    cell.alignment = align or (RT if isinstance(v, (int, float)) else LT)
    cell.border = BD
    if fmt:
        cell.number_format = fmt
    elif isinstance(v, float):
        cell.number_format = "#,##0.00"
    elif isinstance(v, int) and not isinstance(v, bool):
        cell.number_format = "#,##0"
    return cell


def put(ws, r, c, v, font=N, fill=None, align=None, fmt=None, bold=False):
    if isinstance(c, list):          # put(ws, r, [values], [fonts])
        vs, start = c, 1
    else:
        vs, start = (v if isinstance(v, list) else [v]), c
    fs = font if isinstance(font, list) else [font] * len(vs)
    cell = None
    for i, (vv, ff) in enumerate(zip(vs, fs)):
        cell = _put_one(ws, r, start + i, vv, ff, fill, align, fmt, bold)
    return cell


def title(ws, t, sub=""):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    c = ws.cell(1, 1, t)
    c.font = TITLE
    c.fill = FILLT
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
        ws.cell(2, 1, sub).font = SUB
    ws.freeze_panes = "A4"


def hd(ws, r, labels, widths=None):
    for i, v in enumerate(labels, 1):
        c = ws.cell(r, i, v)
        c.font = H
        c.fill = FILLH
        c.alignment = CT
        c.border = BD
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def period_of(s):
    if not s:
        return None
    d = datetime.date.fromisoformat(str(s)[:10])
    y, m = (d.year + 1, 1) if (d.month == 12 and d.day >= 26) else \
           (d.year, d.month + 1) if d.day >= 26 else (d.year, d.month)
    return "%d-%02d期" % (y, m)


def period_range(key):
    """2026-10期 → (2026-09-26, 2026-10-25)"""
    y, m = int(key[:4]), int(key[5:7])
    start = datetime.date(y, m, 26) if m > 1 else datetime.date(y - 1, 12, 26)
    # 期起 = 上月26日
    sy, sm = (y, m - 1) if m > 1 else (y - 1, 12)
    start = datetime.date(sy, sm, 26)
    end = datetime.date(y, m, 25)
    return start, end


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def load_scope(periods):
    conn = sqlite3.connect(f"file:{_db_path()}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    trips = [dict(t) for t in conn.execute("""
        select t.*, ft.short_name chain, ft.code tpl_code,
               d.name dispname, s.name srcname, s.settlement_type, s.rate_point,
               s.outbound_rate_point, s.quota_amount, a.name accname
        from trips t
        left join flow_templates ft on ft.id = t.template_id
        left join dispatchers d on d.id = t.dispatcher_id
        left join invoice_sources s on s.id = t.invoice_source_id
        left join accounts a on a.id = t.account_id
        where t.status != 'voided' order by t.settlement_date, t.id""")]
    srcs = {r["id"]: dict(r) for r in conn.execute("SELECT * FROM invoice_sources")}
    # 票源方「分期额度 / 接单状态」（2026-10-04）：票额按月给、中途停接单这类时效性规则在这里，
    # 报告按"所报最新一期的起始日"解析生效行，未命中则回退 invoice_sources 的档案默认值。
    try:
        terms = [dict(r) for r in conn.execute(
            "SELECT * FROM invoice_source_terms ORDER BY invoice_source_id, effective_from")]
    except sqlite3.OperationalError:
        terms = []          # 老库还没建这张表时不报错，退回档案默认值
    tiers = [dict(r) for r in conn.execute("SELECT * FROM mill_reward_tiers ORDER BY min_tonnage")]
    batches = [dict(b) for b in conn.execute("SELECT * FROM batches ORDER BY id")]
    bt = {}
    for r in conn.execute("SELECT batch_id, trip_id, amount FROM batch_trips"):
        bt[r["trip_id"]] = (r["batch_id"], num(r["amount"]))
    F_KEYS = ("give_customer", "front_invoice_amount", "front_actual_paid", "personal_advance",
              "service_fee_income", "reward_profit", "dispatcher_commission_cost",
              "net_profit", "invoice_diff_profit")
    scope = []
    all_scope = []          # 全库车次（不分期）—— 供「年度额度水位」跨期累计用
    for t in trips:
        p = period_of(t["settlement_date"])
        ov = json.loads(t["finance_overrides"] or "{}")
        t["_p"] = p
        f = engine.compute_trip_finance_v2_flat(conn, t, ov) or {}
        t["_B"] = num(t["settlement_tonnage_total"])
        t["_D"] = num(t["settlement_amount_total"])
        t["_advu"] = num(engine.advance_fee_unit_for_trip(conn, t))
        t["_adv"] = t["_B"] * t["_advu"]
        t["_svc"] = t["tpl_code"] == "service_fee_direct"
        t["_give"] = num(f.get("give_customer"))
        t["_front_inv"] = num(f.get("front_invoice_amount"))
        t["_front_paid"] = num(f.get("front_actual_paid"))
        t["_pa"] = num(f.get("personal_advance"))
        t["_rw"] = num(f.get("reward_profit"))
        t["_rwu"] = num(f.get("reward_unit_price"))
        t["_cm"] = num(f.get("dispatcher_commission_cost"))
        t["_sfee"] = num(f.get("service_fee_income"))
        t["_net"] = num(f.get("net_profit"))
        t["_ires"] = num(f.get("invoice_diff_profit"))
        t["_done"] = bt.get(t["id"], (None, 0.0))[1]
        all_scope.append(t)
        if p in periods:
            scope.append(t)
    conn.close()
    return scope, all_scope, srcs, tiers, batches, terms


def agg(rows, keys=AGG_KEYS):
    return {k: sum(t["_" + k] for t in rows) for k in keys}


def sheet_overview(wb, periods, scope, sel, batches):
    ws = wb.create_sheet("① 总览")
    title(ws, "经营报告总览（%s）" % "、".join(periods),
          "系统引擎实时口径；共 %d 台车。生成日 %s。" % (len(scope), datetime.date.today().isoformat()))
    r = 4
    ws.cell(r, 1, "A. 分期汇总").font = SB; r += 1
    hd(ws, r, ["期", "结算区间", "车次数", "吨位", "含税结算额", "给客户付", "量补", "提成",
               "个人垫资", "净利润", "净利/吨"], [11, 24, 8, 11, 16, 15, 11, 11, 12, 13, 10]); r += 1
    for p in periods:
        rows = sel(p)
        a = agg(rows)
        s, e = period_range(p)
        put(ws, r, [p, "%s ~ %s" % (s, e), len(rows), a["B"], a["D"], a["give"], a["rw"],
                    a["cm"], a["pa"], a["net"], (a["net"] / a["B"]) if a["B"] else 0],
            [None, None, NT, MT, MT, MT, MT, MT, MT, MT, MT],
            bold=True, fill=GRN); r += 1
    a = agg(scope)
    put(ws, r, ["合计", "", len(scope), a["B"], a["D"], a["give"], a["rw"], a["cm"], a["pa"], a["net"],
                (a["net"] / a["B"]) if a["B"] else 0],
        [None, None, NT, MT, MT, MT, MT, MT, MT, MT, MT], bold=True, fill=BLU); r += 2

    ws.cell(r, 1, "B. 分期 × 链条").font = SB; r += 1
    hd(ws, r, ["期", "链条", "车次数", "吨位", "含税结算额", "量补", "提成", "个人垫资", "净利润"],
       [11, 11, 8, 11, 16, 11, 11, 12, 13]); r += 1
    for p in periods:
        for ch in CHAIN_ORDER:
            rows = sel(p, ch)
            if not rows:
                continue
            x = agg(rows)
            put(ws, r, [p, ch, len(rows), x["B"], x["D"], x["rw"], x["cm"], x["pa"], x["net"]],
                [None, None, NT, MT, MT, MT, MT, MT, MT]); r += 1
        x = agg(sel(p))
        put(ws, r, [p, "小计", len(sel(p)), x["B"], x["D"], x["rw"], x["cm"], x["pa"], x["net"]],
            [None, None, NT, MT, MT, MT, MT, MT, MT], bold=True, fill=GRN); r += 1
    r += 1

    ws.cell(r, 1, "C. 关注专题速览").font = SB; r += 1
    heads = ["专题"] + periods + ["合计", "口径说明"]
    hd(ws, r, heads, [22] + [15] * len(periods) + [16, 62]); r += 1
    lkcm = lambda rows: sum(t["_cm"] for t in rows)
    give = lambda rows: sum(t["_D"] - t["_adv"] for t in rows)
    inv = lambda rows: sum(t["_front_inv"] for t in rows)
    TOPICS = [
        ("调货人工资（提成）", lkcm, "按 rate_links 提成档逐车计提（林凯带票15元/吨；邓晶焱带票5元/吨须当期≥1000吨；其余0）"),
        ("量补（钢厂奖励）", lambda rows: sum(t["_rw"] for t in rows),
         "华西特钢档位：≥3000吨→20元/吨、≥5000吨→25元/吨，按期累计整期落档"),
        ("应收·欧冶回款（票额占用）", give, "含税结算额 − 吨位×垫资费率（配票/带票 23 元/吨；反向票源为 0）"),
        ("应付·供货方货款（给客户付）", lambda rows: sum(t["_give"] for t in rows),
         "含税 × (1−14%) − 吨减（非林凯 5 元/吨，2026-09-26 起）；2026-10 期切割费另在税后全额扣"),
        ("前端开票（票源方开票额）", inv, "含税 − 垫资；占用票源方额度，见「票源额度水位」表（A 配票/反向水位、B 自带票流水）"),
    ]
    for lbl, fn, note in TOPICS:
        vals = [fn(sel(p)) for p in periods]
        put(ws, r, [lbl] + vals + [sum(vals), note],
            [None] + [MT] * (len(periods) + 1) + [None], fill=YEL); r += 1
    # 工资已结/未结
    cm_all = lkcm(scope)
    done = sum(t["_done"] for t in scope)
    put(ws, r, ["　└ 已结 / 未结", "", "", "", "", "", "%.2f / %.2f" % (done, cm_all - done),
                "已结 = batches+batch_trips 逐车登记额（跨所选期）"],
        [None] * 6 + [MT, None], fill=GRY); r += 1


def sheet_period(wb, idx, p, scope, sel):
    rows = sel(p)
    a = agg(rows)
    s, e = period_range(p)
    ws = wb.create_sheet("%s %s" % ("②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮"[min(idx, 15) - 1], p))
    title(ws, "%s 经营统计（结算日 %s ~ %s）" % (p, s, e),
          "本期 %d 台车 · %.2f 吨 · 含税 %.2f 元 · 净利润 %.2f 元。全部为系统引擎实时口径。"
          % (len(rows), a["B"], a["D"], a["net"]))
    r = 4
    ws.cell(r, 1, "A. 分链条").font = SB; r += 1
    hd(ws, r, ["链条", "车次数", "吨位", "含税结算额", "给客户付", "前端实付", "量补",
               "个人垫资", "提成", "净利润", "净利/吨"],
       [11, 8, 11, 16, 15, 15, 11, 12, 11, 13, 10]); r += 1
    for ch in CHAIN_ORDER:
        rs = sel(p, ch)
        if not rs:
            continue
        x = agg(rs)
        put(ws, r, [ch, len(rs), x["B"], x["D"], x["give"], x["front_paid"], x["rw"],
                    x["pa"], x["cm"], x["net"], (x["net"] / x["B"]) if x["B"] else 0],
            [None, NT, MT, MT, MT, MT, MT, MT, MT, MT, MT]); r += 1
    x = agg(rows)
    put(ws, r, ["小计", len(rows), x["B"], x["D"], x["give"], x["front_paid"], x["rw"],
                x["pa"], x["cm"], x["net"], (x["net"] / x["B"]) if x["B"] else 0],
        [None, NT, MT, MT, MT, MT, MT, MT, MT, MT, MT], bold=True, fill=GRN); r += 2

    ws.cell(r, 1, "B. 逐车明细").font = SB; r += 1
    hd(ws, r, ["车次#", "结算日", "链条", "调货人", "票源方", "吨位", "含税结算额", "给客户付",
               "前端开票", "前端实付", "量补单价", "量补", "垫资单价", "垫资", "提成", "其他成本", "净利润", "已结工资"],
       [7, 11, 9, 9, 13, 9, 13, 12, 12, 12, 9, 9, 9, 9, 9, 9, 12, 11]); r += 1
    for t in rows:
        oc = num(t.get("other_cost"))
        put(ws, r, ["#%d" % t["id"], str(t["settlement_date"])[:10], t["chain"],
                    t["dispname"] or "-", t["srcname"] or "-", t["_B"], t["_D"], t["_give"],
                    t["_front_inv"], t["_front_paid"], t["_rwu"] or None, t["_rw"],
                    t["_advu"] or None, t["_adv"], t["_cm"], oc, t["_net"], t["_done"]],
            [None, None, None, None, None] + [MT] * 9 + [MT, MT, MT, None, MT, MT],
            fill=YEL if t["dispname"] == "林凯" else GRY); r += 1
    x = agg(rows)
    put(ws, r, ["合计", "", "", "", "", x["B"], x["D"], x["give"], x["front_inv"], x["front_paid"],
                "", x["rw"], "", x["adv"], x["cm"], "", x["net"], sum(t["_done"] for t in rows)],
        [None] * 5 + [MT] * 5 + [None, MT, None, MT, MT, None, MT, MT],
        bold=True, fill=BLU)
    ws.freeze_panes = "A%d" % (r - len(rows))
    return ws


def sheet_wages(wb, periods, scope, sel):
    ws = wb.create_sheet("调货人工资")
    title(ws, "调货人工资（提成）—— 分人 × 路径 × 期",
          "按 rate_links 提成档逐车计提；已结 = 系统批次登记额。")
    r = 4
    hd(ws, r, ["调货人", "路径", "车次数", "吨位", "应结工资"] + periods + ["已结", "未结"],
       [11, 11, 8, 11, 13] + [13] * len(periods) + [12, 12]); r += 1
    combos = sorted({(t["dispname"] or "(未填)", t["chain"]) for t in scope},
                    key=lambda x: (x[0], CHAIN_ORDER.index(x[1]) if x[1] in CHAIN_ORDER else 9))
    for nm, ch in combos:
        rows = [t for t in scope if (t["dispname"] or "(未填)") == nm and t["chain"] == ch]
        if not rows:
            continue
        cm = sum(t["_cm"] for t in rows)
        per = [sum(t["_cm"] for t in rows if t["_p"] == p) for p in periods]
        done = sum(t["_done"] for t in rows)
        put(ws, r, [nm, ch, len(rows), sum(t["_B"] for t in rows), cm] + per + [done, cm - done],
            [None, None, NT, MT, MT] + [MT] * len(periods) + [MT, MT],
            bold=(nm == "林凯"), fill=GRN if nm == "林凯" else GRY); r += 1
    cm_all = sum(t["_cm"] for t in scope)
    put(ws, r, ["合计", "", len(scope), sum(t["_B"] for t in scope), cm_all]
        + [sum(t["_cm"] for t in scope if t["_p"] == p) for p in periods]
        + [sum(t["_done"] for t in scope), cm_all - sum(t["_done"] for t in scope)],
        [None, None, NT, MT, MT] + [MT] * len(periods) + [MT, MT], bold=True, fill=BLU); r += 2
    ws.cell(r, 1, "未结口径：应结 − 系统批次已登记。邓晶焱带票 5 元/吨须当期带票累计 ≥1000 吨（整期口径），未达档按 0。").font = SUB


def sheet_reward(wb, periods, scope, sel, tiers):
    ws = wb.create_sheet("量补汇总")
    title(ws, "量补（钢厂奖励）—— 分期 × 档位 × 计算过程",
          "量补 = 该期累计结算吨位落档单价 × 整期吨位，逐车同价。来源 mill_reward_tiers。")
    r = 4
    hd(ws, r, ["期", "钢厂", "车次数", "吨位", "档位单价(元/吨)", "量补金额", "计算过程", "档位来源"],
       [11, 11, 8, 12, 13, 13, 34, 52]); r += 1
    for p in periods:
        rows = sel(p)
        ton = sum(t["_B"] for t in rows)
        rw = sum(t["_rw"] for t in rows)
        us = sorted({round(t["_rwu"], 4) for t in rows if t["_rwu"]})
        unit = us[0] if len(us) == 1 else 0
        src = next(("mill_reward_tiers#%d（≥%s吨→%s元/吨）" % (x["id"], x["min_tonnage"], x["unit_price"])
                    for x in tiers if x["unit_price"] == unit and unit), "混合/无")
        put(ws, r, [p, "华西特钢", len(rows), ton, unit, rw,
                    "%.2f 吨 × %g 元/吨 = %.2f" % (ton, unit, rw), src],
            [None, None, NT, MT, None, MT, None, None], bold=True, fill=GRN if rw else GRY); r += 1
    tot = sum(t["_rw"] for t in scope)
    put(ws, r, ["合计", "华西特钢", len(scope), sum(t["_B"] for t in scope), "", tot, "", ""],
        [None, None, NT, MT, None, MT, None, None], bold=True, fill=BLU); r += 2
    ws.cell(r, 1, "档位参数：%s" % "；".join("≥%s吨→%s元/吨（周期%d日~%d日）"
            % (x["min_tonnage"], x["unit_price"], x["period_start_day"], x["period_end_day"]) for x in tiers)).font = SUB


def _period_start_of(p, start_day=26):
    """'2026-10期' → 该期起始日 '2026-09-26'（26 日 ~ 次月 25 日）。"""
    y, m = int(p[:4]), int(p[5:7])
    if m == 1:
        return "%d-12-%02d" % (y - 1, start_day)
    return "%d-%02d-%02d" % (y, m - 1, start_day)


def _term_for(terms, sid, on_date):
    """该票源方在 on_date 生效的那条分期记录（没有=None，走档案默认值）。"""
    hit = [t for t in terms if t["invoice_source_id"] == sid
           and str(t["effective_from"]) <= on_date
           and (not t["effective_to"] or str(t["effective_to"]) >= on_date)]
    hit.sort(key=lambda t: (str(t["effective_from"]), t["id"]))
    return hit[-1] if hit else None


def _role_of(src, modes):
    """角色由车次结算模式推断：M1 配票 / M5 反向（占额度）；M2~M4 自带票（不占额度）。
    不再用"档案里填没填额度"区分——否则某家"本期没有票额"会被误判成自带票混进 B 表。"""
    if src["source_type"] == "reverse_individual" or "M5" in modes:
        return "反向"
    if "M1" in modes:
        return "配票"
    if modes & {"M2", "M3", "M4"}:
        return "自带票"
    return "配票" if num(src.get("quota_amount")) else "自带票"


def sheet_quota(wb, periods, all_scope, srcs, terms):
    """票源额度水位：A 配票 + 反向（占额度，再按配票/反向分两组） / B 自带票（不占额度，仅统计票面流水）。
    时效：票额按月给、中途停接单 → 按所报最新一期起始日解析 invoice_source_terms；
          被停用或该期无票额的票源不计入水位，只在下方以一行"本期不计入"说明。"""
    ws = wb.create_sheet("票源额度水位")
    latest = sorted(periods)[-1]
    year = latest[:4]
    on_date = _period_start_of(latest)   # 分期规则的判定基准日 = 最新一期起始日（报告可复现，不随生成日漂移）
    title(ws, "票源额度水位 —— 按所选期",
          "A 配票/反向：占用票源方开票额度，水位 = 周期内已用开票额 ÷ 额度"
          "（月度额度看 %s期，年度额度看 %s 年累计）；B 自带票（直签带票 M2）：由带票公司自行开票，"
          "不占我方票源额度，仅统计票面流水、本期优先（往期仅供对比）。前端开票 = 含税 − 吨位×垫资费率（配票 23 元/吨、反向票源 0）。"
          "额度与是否接单按「票源额度分期」表在 %s 的生效值取，该期没票额/已停接单的不计入。"
          % (latest, year, on_date))
    n = len(periods)
    cur_i = periods.index(latest)        # 本期在 periods 中的下标（不依赖升序假设）
    r = 4

    # 车次模式分布 → 角色
    modes = {}
    for t in all_scope:
        if t["invoice_source_id"] and t["mode_code"]:
            modes.setdefault(t["invoice_source_id"], set()).add(t["mode_code"])

    # ═════════ A. 配票 + 反向（占票源额度） ═════════
    ws.cell(r, 1, "A. 配票 + 反向 —— 票源额度水位（占额度）").font = SB
    r += 1
    hd(ws, r, ["票源方", "链条", "水位口径", "额度(元)"] + ["开票·" + p for p in periods] +
              ["周期内已用", "水位", "剩余额度", "状态"],
       [14, 7, 11, 15] + [14] * n + [15, 8, 15, 12])
    r += 1
    a_first = r
    rows_a, skipped = [], []
    for sid, s in sorted(srcs.items()):
        chain = _role_of(s, modes.get(sid, set()))
        if chain == "自带票":
            continue                      # 直签带票（无额度档案）→ B 部分
        # —— 时效：先看分期表在该期起始日命中的那条 ——
        term = _term_for(terms, sid, on_date)
        if term is not None:
            if not (term["accepting"] or 0):
                skipped.append((s["name"], "已停止接单"))
                continue
            if term["quota_amount"] is None:
                skipped.append((s["name"], "本期无票额"))
                continue
            quota = num(term["quota_amount"])
            warn = num(term["warn_threshold"]) if term["warn_threshold"] is not None else num(s.get("warn_threshold"))
        else:
            quota = num(s.get("quota_amount"))
            warn = num(s.get("warn_threshold"))
        own = [t for t in all_scope if t["invoice_source_id"] == sid]
        per = [sum(t["_front_inv"] for t in own if t["_p"] == p) for p in periods]
        qp = s.get("quota_period")
        if qp == "annual":
            used = sum(t["_front_inv"] for t in own if str(t["settlement_date"] or "")[:4] == year)
            qlbl = "年度·累计"
        else:
            used = per[cur_i]
            qlbl = "月度·本期"
        occ = (used / quota) if quota else None
        st = ("○ 本期无用量" if not used else
              "—" if occ is None else
              "⚠ 超额度" if occ >= 1 else
              "⚠ 已过预警线" if (warn and used >= warn) else
              "⚠ 水位偏高" if occ >= _high_water_pct() else "✓ 正常")
        rows_a.append(dict(name=s["name"], chain=chain, qlbl=qlbl, quota=quota, per=per,
                           used=used, occ=occ, rest=(max(quota - used, 0) if quota else None), st=st))
    # 先按链条分两组（配票在前、反向在后），组内用量大的在前；组内插一条小标题行
    chain_tot = {}
    all_per = [0.0] * n
    bars = []          # (数据行号) 用于给水位列加水位数据条
    for chain in ("配票", "反向"):
        grp = sorted([x for x in rows_a if x["chain"] == chain], key=lambda x: -x["used"])
        if not grp:
            continue
        put(ws, r, ["A%d %s（%s）" % (len(chain_tot) + 1, chain,
                    "前端公司 · 额度按月" if chain == "配票" else "自然人 · 额度按年")],
            [Font(name="微软雅黑", size=10, bold=True, color="1F4E79")], fill=FILLH)
        r += 1
        tp, tq, tu = [0.0] * n, 0.0, 0.0
        for x in grp:
            put(ws, r, [x["name"], x["chain"], x["qlbl"], x["quota"]] + x["per"] +
                       [x["used"], x["occ"], x["rest"], x["st"]],
                [H, None, None, M0] + [MT] * n + [MT, MT, MT, None],
                fill=YEL if x["used"] else GRY)
            if x["occ"] is not None:
                ws.cell(r, 6 + n).number_format = "0.0%"
            bars.append(r)
            for i, v in enumerate(x["per"]):
                tp[i] += v
                all_per[i] += v
            tq += x["quota"]
            tu += x["used"]
            r += 1
        put(ws, r, ["小计", chain, "月度/年度混合", tq] + tp +
                   [tu, (tu / tq) if tq else None, max(tq - tu, 0), ""],
            [H, None, None, M0] + [MT] * n + [MT, MT, MT, None], bold=True, fill=BLU)
        if tq:
            ws.cell(r, 6 + n).number_format = "0.0%"
        chain_tot[chain] = (tq, tu)
        r += 1
    if bars:
        col = get_column_letter(6 + n)
        for rr in bars:                 # 水位列加水位数据条（逐行，跳过小标题/小计行）
            ws.conditional_formatting.add(
                "%s%d" % (col, rr),
                DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                            color="638EC6", showValue=True))
    if chain_tot:
        tq = sum(v[0] for v in chain_tot.values())
        tu = sum(v[1] for v in chain_tot.values())
        put(ws, r, ["合计", "配票+反向", "月度/年度混合", tq] + all_per +
                   [tu, (tu / tq) if tq else None, max(tq - tu, 0), ""],
            [H, None, None, M0] + [MT] * n + [MT, MT, MT, None], bold=True, fill=BLU)
        if tq:
            ws.cell(r, 6 + n).number_format = "0.0%"
        r += 1
    if skipped:
        put(ws, r, ["本期不计入：" + "、".join("%s（%s）" % (nm, why) for nm, why in skipped)],
            [SM], fill=GRY)
        r += 1
    if not chain_tot:
        put(ws, r, ["（所选期内没有占用额度的票源方）"], [SM], fill=GRY)
        r += 1

    r += 1
    # ═════════ B. 自带票（直签带票 M2，不占额度）—— 本期优先 ═════════
    ws.cell(r, 1, "B. 自带票（直签带票 M2）—— 不占我方额度，只看本期多少（本期 = %s）" % latest).font = SB
    r += 1
    past = [p for p in periods if p != latest]
    hd(ws, r, ["开票公司", "本期车次", "本期票面额（%s）" % latest, "本期吨位"] +
              ["票面额·" + p for p in past] + ["累计票面额（含往期）", "累计吨位", "备注"],
       [15, 9, 19, 12] + [14] * len(past) + [19, 12, 30])
    r += 1
    b_start = r
    tot_cur, tot_cur_ton, tot_past, tn_cur, cum_all, cum_ton = 0.0, 0.0, [0.0] * len(past), 0, 0.0, 0.0
    for sid, s in sorted(srcs.items()):
        if _role_of(s, modes.get(sid, set())) != "自带票":
            continue
        own = [t for t in all_scope if t["invoice_source_id"] == sid and t["_p"] in periods]
        if not own:
            continue
        own_cur = [t for t in own if t["_p"] == latest]
        per = [sum(t["_D"] for t in own if t["_p"] == p) for p in past]
        cur = sum(t["_D"] for t in own_cur)
        cur_ton = sum(t["_B"] for t in own_cur)
        put(ws, r, [s["name"], len(own_cur), cur, cur_ton] + per +
                   [sum(per) + cur, sum(t["_B"] for t in own), ""],
            [H, None, MT, MT] + [MT] * len(past) + [MT, MT, None],
            fill=YEL if cur else GRY)
        # 本期两列加粗，突出「这个月多少」
        for cc in (3, 4):
            ws.cell(r, cc).font = Font(name="微软雅黑", size=10, bold=True)
        tot_cur += cur
        tot_cur_ton += cur_ton
        for i, v in enumerate(per):
            tot_past[i] += v
        tn_cur += len(own_cur)
        cum_all += sum(per) + cur
        cum_ton += sum(t["_B"] for t in own)
        r += 1
    if r > b_start:
        put(ws, r, ["合计", tn_cur, tot_cur, tot_cur_ton] + tot_past +
                   [cum_all, cum_ton, "票面额 = 含税结算额（公司自行开票）；本期即 %s" % latest],
            [H, None, MT, MT] + [MT] * len(past) + [MT, MT, None], bold=True, fill=BLU)
        r += 1
    else:
        put(ws, r, ["（所选期内无自带票车次）"], [SM], fill=GRY)
        r += 1
    r += 1
    ws.cell(r, 1, "注：A 表额度取自「票源额度分期」(invoice_source_terms) 在 %s 的生效值，未设分期的回退票源方档案 "
                  "(invoice_sources.quota_amount)；配票公司为月度额度、反向个人为年度额度。链条（配票/反向/自带票）由车次结算模式 "
                  "M1/M5/M2~M4 判定，不再看「有没有额度」。「开票·某期」列一律为该期实际前端开票额，便于跨期对比。"
                  "B 表仅供对账，不含额度限制，左起「本期车次 / 本期票面额 / 本期吨位」即 %s 的量，往期列仅作对比。" % (on_date, latest)).font = SUB


def discover_periods():
    conn = sqlite3.connect(f"file:{_db_path()}?mode=ro", uri=True)
    keys_map = {}
    for (sd,) in conn.execute(
            "SELECT settlement_date FROM trips WHERE status!='voided' AND settlement_date IS NOT NULL"):
        k = period_of(sd)
        if k:
            keys_map[k] = True
    conn.close()
    return sorted(keys_map)   # period_of 已带「期」后缀


def build_workbook(periods_param=None):
    """生成工作簿。periods_param：None=自动全部期 / "2026-08,2026-10" / 列表。
    返回 (wb, periods, scope, sel)。库中无车次或期号不合法时抛 ValueError。"""
    keys = discover_periods()
    if not keys:
        raise ValueError("库中无任何车次")
    if periods_param:
        req = periods_param if isinstance(periods_param, (list, tuple)) else \
            [x.strip() for x in str(periods_param).split(",") if x.strip()]
        want = ["%s期" % x if "期" not in x else x for x in req]
        bad = [w for w in want if w not in keys]
        if bad:
            raise ValueError("库中无这些期（已有：%s）" % ", ".join(sorted(keys)))
        periods = want
    else:
        periods = keys

    scope, all_scope, srcs, tiers, batches, terms = load_scope(set(periods))

    def sel(p=None, chain=None):
        r = [t for t in scope if p is None or t["_p"] == p]
        return [x for x in r if x["chain"] == chain] if chain else r

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_overview(wb, periods, scope, sel, batches)
    for i, p in enumerate(periods, 1):
        sheet_period(wb, i, p, scope, sel)
    sheet_wages(wb, periods, scope, sel)
    sheet_reward(wb, periods, scope, sel, tiers)
    sheet_quota(wb, periods, all_scope, srcs, terms)
    return wb, periods, scope, sel


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--periods", help="逗号分隔，如 2026-08,2026-10；缺省=自动探测全部期")
    ap.add_argument("--out", help="输出文件名（不含路径，缺省写到 data/reports/）")
    args = ap.parse_args()

    try:
        wb, periods, scope, sel = build_workbook(args.periods)
    except ValueError as e:
        print("⚠", e)
        raise SystemExit(1)

    tag = "_".join(x.replace("期", "") for x in periods)
    out = args.out or "经营报告_%s_%s.xlsx" % (tag, datetime.date.today().isoformat())
    path = out if os.path.isabs(out) or os.path.dirname(out) else os.path.join(REPORTS_DIR, out)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print("已生成:", path)
    for p in periods:
        rows = sel(p)
        a = agg(rows)
        print("  %s: %d 台 %.2f 吨 净利 %.2f 量补 %.2f 提成 %.2f" % (p, len(rows), a["B"], a["net"], a["rw"], a["cm"]))


if __name__ == "__main__":
    main()
