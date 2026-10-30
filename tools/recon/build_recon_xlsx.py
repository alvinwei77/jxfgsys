import json, re, datetime as dt, os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

"""生成分步对账 Excel（总览/①车次/②调货人供货厂/③票源/④计算路径/⑤计算过程/⑥路径统计/⑦各路径计算过程/⑧同步明细）。
用法：python3 build_recon_xlsx.py match.json 系统车次.json 输出.xlsx [同步前系统车次.json 同步改动.json]
没有同步前数据时，⑧和“本次同步改了什么”列为空。注意：总览里的文字结论是按 2026-09-22 那次对账写的，换数据后要核对并改写。"""
import sys
M = json.load(open(sys.argv[1]))
S = {s['id']: {**s, 'plate': s['plate'] or '（车号待补）'} for s in json.load(open(sys.argv[2]))}
OUT = sys.argv[3]
SB = {s['id']: s for s in json.load(open(sys.argv[4]))} if len(sys.argv) > 4 else S
CHG = {c['id']: c for c in json.load(open(sys.argv[5]))} if len(sys.argv) > 5 else {}
CHG = {i: CHG.get(i, {'changes': []}) for i in S}

FNT = 'Arial'
thin = Side(style='thin', color='D0D5DD'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill('solid', fgColor='1F3A68'); HT = Font(name=FNT, bold=True, color='FFFFFF', size=10)
HF2 = PatternFill('solid', fgColor='41557A')
F = Font(name=FNT, size=10); FB = Font(name=FNT, size=10, bold=True); FG = Font(name=FNT, size=9, color='667085')
FGR = Font(name=FNT, size=10, color='98A2B3')
C_RED = PatternFill('solid', fgColor='F8B4B4'); C_ORG = PatternFill('solid', fgColor='FDD9A8')
C_BLU = PatternFill('solid', fgColor='CFE3FA'); C_GRN = PatternFill('solid', fgColor='D6F0DD')
C_INT = PatternFill('solid', fgColor='EEF0F3'); C_GRY = PatternFill('solid', fgColor='F2F4F7'); INP = PatternFill('solid', fgColor='FFFBEA')
C_FIN = PatternFill('solid', fgColor='FFF7E6'); C_SYS = PatternFill('solid', fgColor='EAF2FF')
MONEY = '#,##0.00;[Red]-#,##0.00;-'
TPL2PATH = {'配票': '配票', '反向': '反向', '直签带票': '带票', '户头通道': '带票', '调货对接': '带票'}
NORM = {'闫东鹏': '闫东朋', '鑫景通': '鑫景通'}
nm = lambda x: NORM.get(x, x)
PEOPLE = set(n for s in S.values() for n in (s['disp'], s['base']) if n)


def parse_j(jf):
    s = jf.lstrip('=').replace(' ', '')
    k = base = None
    m = re.search(r'([ED])\d+\*(0\.\d+)', s)
    if m: base, k = m.group(1), float(m.group(2))
    elif re.search(r'E\d+-F\d+', s): base, k = 'E', 1.0
    r = c = 0.0
    for sign, a, b in re.findall(r'([+-]?)(?:B\d+\*(\d+)|(\d+)\*B\d+)', s):
        n = float(a or b)
        if sign == '-': c += n
        else: r += n
    return k, base, r, c


def header(ws, hdr, row=1, widths=None, h=32, fills=None):
    for i, t in enumerate(hdr, 1):
        c = ws.cell(row, i, t); c.font = HT; c.fill = (fills or {}).get(i, HF); c.border = B
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = h
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[CL(i)].width = w


def style_row(ws, r, n, money=(), num=(), wrap=()):
    for j in range(1, n + 1):
        c = ws.cell(r, j); c.font = F; c.border = B; c.alignment = Alignment(vertical='top', wrap_text=j in wrap)
        if j in money: c.number_format = MONEY
        if j in num: c.number_format = '0.00'


def paint(ws, r, cols, fill):
    for c in cols: ws.cell(r, c).fill = fill


# ---------- 组装财务行 ----------
R = []
for o in M['rows']:
    f, s = o['f'], (S[o['s']['id']] if o['s'] else None)
    sh = f['sheet']
    if sh.startswith('林凯-'): fd, fb = '林凯', sh.split('-', 1)[1]
    elif sh == '林凯': fd, fb = '林凯', None
    else: fd, fb = sh, None
    fsrc = '自带票' if f['path'] == '带票' else ('反向' if f['path'] == '反向' else (f['src'] or '（财务没写）'))
    k, base, rr, cc = parse_j(f['jf'])
    x = dict(f=f, s=s, fd=fd, fb=fb, fsrc=fsrc, k=k, kbase=base, r_f=rr, c_f=cc, merged=False, missing=False)
    if s is None:  # 财务一行 = 系统两车？（同日期±1天、吨位和金额合计都对得上）
        pair = None
        un = M.get('unmatched_sys', [])
        for i in range(len(un)):
            for j in range(i + 1, len(un)):
                a, b = un[i], un[j]
                if abs(a['ton'] + b['ton'] - f['ton']) < 0.01 and abs(a['amt'] + b['amt'] - f['amt']) < 1:
                    pair = [S[a['id']], S[b['id']]]
        if pair: x.update(merged=True, subs=pair)
        else: x['missing'] = True
    R.append(x)

wb = Workbook()
INTNOTE = 'INT'

# ====================== ① 车次对账 ======================
w1 = wb.active; w1.title = '①车次对账'
header(w1, ['财务表', '行', '财务日期', '财务吨位', '财务应收金额', '系统车次', '车号', '系统结算日期', '系统吨位', '系统结算金额', '对账结果', '说明'],
       widths=[11, 5, 10, 9, 13, 9, 20, 11, 9, 13, 14, 60], fills={i: HF2 for i in range(6, 11)})
r = 2
for x in R:
    f = x['f']
    if x['missing']:
        vals = [f['sheet'], f['row'], f['date'], f['ton'], f['amt'], '', '', '', '', '', '系统未录', '9/21 闫东鹏反向，你还没录入（录入后利润应≈840.59）']
    elif x['merged']:
        a, b = x['subs']
        vals = [f['sheet'], f['row'], f['date'], f['ton'], f['amt'], '#33+#34', f"{a['plate']} / {b['plate']}", a['sd'], a['ton'] + b['ton'], a['amt'] + b['amt'], '合并',
                f"财务一行 = 系统两车：{a['plate']} {a['ton']}吨 + {b['plate']} {b['ton']}吨，吨位金额合计一致"]
    else:
        s = x['s']; notes = []; res = '一致'
        if abs(s['ton'] - f['ton']) > 0.005: notes.append(f"吨位：财务{f['ton']}，系统{s['ton']}"); res = '不一致'
        if abs(s['amt'] - f['amt']) > 1: notes.append(f"金额：财务{f['amt']:,.2f}，系统{s['amt']:,.2f}"); res = '不一致'
        if s['sd'] != f['date']:
            dd = abs((dt.date.fromisoformat(s['sd']) - dt.date.fromisoformat(f['date'])).days)
            notes.append(f"日期差{dd}天" + ('（船，财务是结算单日期）' if s['type'] == 'ship' else ''))
            if res == '一致': res = '一致(日期差)'
        vals = [f['sheet'], f['row'], f['date'], f['ton'], f['amt'], f"#{s['id']}", s['plate'], s['sd'], s['ton'], s['amt'], res, '；'.join(notes)]
    for j, v in enumerate(vals, 1): w1.cell(r, j, v)
    style_row(w1, r, 12, money=(5, 10), num=(4, 9), wrap=(12,))
    res = vals[10]
    w1.cell(r, 11).fill = {'一致': C_GRN, '一致(日期差)': C_BLU, '合并': C_ORG}.get(res, C_RED)
    if x['missing']: paint(w1, r, range(1, 11), C_RED)
    if x['merged']: paint(w1, r, (6, 7, 9, 10), C_ORG)
    if not x['missing'] and not x['merged'] and x['s']['sd'] != f['date']: paint(w1, r, (3, 8), C_BLU)
    r += 1
L1 = r - 1; w1.freeze_panes = 'F2'; w1.auto_filter.ref = f'A1:L{L1}'


# ====================== ② 调货人/供货厂 ======================
def attr(x, s):
    fd, fb = x['fd'], x['fb']; sd, sb = s['disp'], s['base']; notes = []
    sheet = x['f']['sheet']
    if sheet != '林凯' and not sheet.startswith('林凯-'):
        p = nm(fd)
        if p in (sd, sb):
            if fd != p: notes.append(f'名字写法：财务“{fd}”，系统“{p}”'); return '写法不同', notes
            return '一致', notes
        return '不一致', [f'财务记在“{fd}”名下，系统是 {sd or "无调货人"}：{sb or "-"}']
    res = '一致'
    if sd != '林凯': res = '不一致'; notes.append(f'调货人：财务林凯，系统{sd or "（未填）"}')
    if fb and nm(fb) != sb: res = '不一致'; notes.append(f'供货厂：财务{fb}，系统{sb or "（未填）"}')
    return res, notes


w2 = wb.create_sheet('②调货人供货厂')
header(w2, ['财务表', '行', '财务日期', '车号', '系统车次', '财务：调货人', '财务：供货厂', '系统调货人', '系统供货厂', '对比结果', '说明', '本次同步改了什么'],
       widths=[11, 5, 10, 20, 9, 14, 12, 11, 11, 12, 44, 40], fills={8: HF2, 9: HF2})
r = 2
for x in R:
    if x['missing']: continue
    subs = x['subs'] if x['merged'] else [x['s']]
    for s in subs:
        res, notes = attr(x, s)
        ch = [c for c in CHG[s['id']]['changes'] if c[0] in ('调货人', '供货厂')]
        chs = '；'.join(f"{a}：{b} → {c}" for a, b, c in ch)
        person = x['f']['sheet'] != '林凯' and not x['f']['sheet'].startswith('林凯-')
        vals = [x['f']['sheet'], x['f']['row'], x['f']['date'], s['plate'], f"#{s['id']}", x['fd'] + ('（按人分表）' if person else ''),
                x['fb'] or ('—' if x['f']['sheet'] == '林凯' else ''), s['disp'] or '（未填）', s['base'] or '（未填）', res,
                '；'.join(notes) + ('；财务这一行是两车合并' if x['merged'] else ''), chs]
        for j, v in enumerate(vals, 1): w2.cell(r, j, v)
        style_row(w2, r, 12, wrap=(11, 12))
        w2.cell(r, 10).fill = {'一致': C_GRN, '写法不同': C_BLU}.get(res, C_RED)
        if res == '写法不同': paint(w2, r, (6, 8), C_BLU)
        if chs: w2.cell(r, 12).fill = C_ORG
        r += 1
L2 = r - 1; w2.freeze_panes = 'E2'; w2.auto_filter.ref = f'A1:L{L2}'
w2.cell(L2 + 2, 1, '“按人分表”的李涛/邓晶焱/高栾凤/闫东鹏/尹鹏飞/王荣胜/陈杰表：系统里调货人或供货厂任一是这个人就算一致。“林凯”表财务没写供货厂。橙色=本次按财务同步改过的。').font = FG

# ====================== ③ 票源 ======================
w3 = wb.create_sheet('③票源')
header(w3, ['财务表', '行', '财务日期', '车号', '系统车次', '财务票源', '系统票源', '对比结果', '说明', '本次同步改了什么'],
       widths=[11, 5, 10, 20, 9, 11, 13, 12, 44, 30], fills={7: HF2})
r = 2
for x in R:
    if x['missing']: continue
    subs = x['subs'] if x['merged'] else [x['s']]
    for s in subs:
        ss = s['src'] or '自带票'; ssn = '反向' if ss.startswith('反向') else ss
        fs = x['fsrc']; note = ''
        if fs == '（财务没写）': res = '财务没写'; note = f'财务这一行没写票源，系统是{ss}（利润按{ss}算，和财务利润一致）'
        elif fs == ssn:
            res = '一致'
            if ss.startswith('反向') and x['f']['sheet'] in ('闫东鹏', '尹鹏飞'):
                exp = '反向-闫东朋' if x['f']['sheet'] == '闫东鹏' else '反向-尹鹏飞'
                if ss != exp: res = '不一致'; note = f'反向个人应为{exp}'
        else: res = '不一致'; note = f'财务{fs}，系统{ss}'
        ch = [c for c in CHG[s['id']]['changes'] if c[0] == '票源']
        chs = '；'.join(f"{b} → {c}" for a, b, c in ch)
        vals = [x['f']['sheet'], x['f']['row'], x['f']['date'], s['plate'], f"#{s['id']}", fs, ss, res, note, chs]
        for j, v in enumerate(vals, 1): w3.cell(r, j, v)
        style_row(w3, r, 10, wrap=(9, 10))
        w3.cell(r, 8).fill = {'一致': C_GRN, '财务没写': C_BLU}.get(res, C_RED)
        if chs: w3.cell(r, 10).fill = C_ORG
        r += 1
L3 = r - 1; w3.freeze_panes = 'E2'; w3.auto_filter.ref = f'A1:J{L3}'

# ====================== ④ 计算路径 ======================
w4 = wb.create_sheet('④计算路径')
header(w4, ['财务表', '行', '财务日期', '车号', '系统车次', '财务路径', '系统链条规则', '对比结果', '本次同步改了什么'],
       widths=[11, 5, 10, 20, 9, 9, 12, 12, 34], fills={7: HF2})
r = 2
for x in R:
    if x['missing']: continue
    subs = x['subs'] if x['merged'] else [x['s']]
    for s in subs:
        sp = TPL2PATH.get(s['tpl'], s['tpl']); res = '一致' if sp == x['f']['path'] else '不一致'
        ch = [c for c in CHG[s['id']]['changes'] if c[0] == '链条规则']
        chs = '；'.join(f"{b} → {c}" for a, b, c in ch)
        vals = [x['f']['sheet'], x['f']['row'], x['f']['date'], s['plate'], f"#{s['id']}", x['f']['path'], s['tpl'], res, chs]
        for j, v in enumerate(vals, 1): w4.cell(r, j, v)
        style_row(w4, r, 9, wrap=(9,))
        w4.cell(r, 8).fill = C_GRN if res == '一致' else C_RED
        if chs: w4.cell(r, 9).fill = C_ORG
        r += 1
L4 = r - 1; w4.freeze_panes = 'E2'; w4.auto_filter.ref = f'A1:I{L4}'


# ====================== ⑤ 计算过程对比 ======================
def sys_steps(subs):
    """把系统一车(或合并的两车)拆成和财务表同样的步骤。"""
    agg = defaultdict(float)
    for s in subs:
        t = s['ton']; A = s['amt']
        agg['ton'] += t; agg['A'] += A
        if s['fam'] == 'split':
            agg['adv'] += s['adv_raw'] or 0; agg['E'] += s['front'] or 0; agg['F'] += s['give'] or 0
            agg['G'] += s['front_paid'] or 0; agg['P'] += s['padv'] or 0
        else:
            agg['adv'] += s['adv'] or 0; agg['E'] += A - (s['adv'] or 0)
            agg['F'] += A - (s['svc'] or 0); agg['G'] += s['svc'] or 0
        agg['rw'] += s['reward'] or 0; agg['cm'] += s['comm'] or 0; agg['oth'] += s['other'] or 0; agg['net'] += s['net'] or 0
        agg['cut'] += s['cut'] or 0
    return agg


def fin_steps(x):
    f = x['f']; t = f['ton']; A = f['amt']; path = f['path']
    d = dict(ton=t, A=A)
    if path == '带票':
        d['adv'] = 23 * t; d['E'] = A - 23 * t; d['F'] = f['pay']; d['G'] = A - f['pay']
    else:
        d['adv'] = 0 if path == '反向' else 23 * t
        d['E'] = A - d['adv']; d['F'] = f['pay']
        d['G'] = (A if x['kbase'] == 'D' else d['E']) * x['k']
        d['P'] = d['F'] - d['G']
    d['rw'] = x['r_f'] * t; d['cm'] = x['c_f'] * t; d['net'] = f['profit']
    return d


w5 = wb.create_sheet('⑤计算过程对比')
w5['A1'] = ('逐车把计算过程拆开对比：左边橙底=财务表，右边蓝底=系统。配票/反向：结算额 → 垫资费 → 前端开票额(欧冶) → 给客户付(实际付款) → 前端实付(带票公司付) → 个人垫资 → 量补、提成 → 利润。'
            '带票：结算额 → 垫资 → 欧冶 → 付供货厂 → 毛服务费(40/吨) → 量补、提成 → 利润。差额>1元标红；0.01~1元标灰“INT”(财务实际付款用了 INT 向下取整，可忽略)。')
w5['A1'].font = FG; w5.merge_cells('A1:AH1'); w5.row_dimensions[1].height = 42; w5['A1'].alignment = Alignment(wrap_text=True, vertical='top')
STEPS = [('A', '结算额'), ('adv', '垫资费'), ('E', '前端开票额\n(欧冶)'), ('F', '给客户付/\n付供货厂'), ('G', '前端实付/\n毛服务费'), ('P', '个人垫资'), ('rw', '量补'), ('cm', '提成'), ('net', '利润')]
hdr = ['财务表', '行', '车号', '系统车次', '路径', '票源', '吨位', '点位(财务→系统)']
fills = {}
for key, lab in STEPS:
    hdr += [f'{lab}\n财务', f'{lab}\n系统', '差']
for i in range(9, 9 + 3 * len(STEPS), 3): fills[i] = PatternFill('solid', fgColor='8A5A00'); fills[i + 1] = HF2; fills[i + 2] = PatternFill('solid', fgColor='667085')
hdr += ['说明']
header(w5, hdr, row=2, widths=[10, 5, 17, 9, 6, 9, 8, 16] + [11, 11, 7] * len(STEPS) + [40], h=44, fills=fills)
r = 3
for x in R:
    if x['missing']: continue
    subs = x['subs'] if x['merged'] else [x['s']]
    fd = fin_steps(x); sd = sys_steps(subs)
    s0 = subs[0]
    if x['f']['path'] == '带票': pt = '付供货厂 结算额−40/吨'
    else:
        gs = (1 - (s0['rate_out'] or 14) / 100); ks = 1 - (s0['rate_in'] or 0) / 100
        gf = round(fd['F'] / fd['A'], 4) if fd['A'] else None
        pt = f"给客户 ×{gf:.3f}→×{gs:.3f}；实付 ×{x['k']}→×{ks:.3f}"
    vals = [x['f']['sheet'], x['f']['row'], ' / '.join(s['plate'] for s in subs), '+'.join(f"#{s['id']}" for s in subs), x['f']['path'],
            ('自带票' if x['f']['path'] == '带票' else (s0['src'] or '')), sd['ton'], pt]
    for j, v in enumerate(vals, 1): w5.cell(r, j, v)
    col = 9; notes = []
    for key, lab in STEPS:
        fv = fd.get(key); sv = sd.get(key) if (key != 'P' or x['f']['path'] != '带票') else None
        w5.cell(r, col, fv); w5.cell(r, col + 1, sv)
        if fv is not None and sv is not None:
            w5.cell(r, col + 2, f'={CL(col+1)}{r}-{CL(col)}{r}')
        col += 3
    if sd['cut']: notes.append(f"切割费{sd['cut']:.0f}：两边结算额里都已扣，客户承担，利润不再扣（已统一）")
    if x['merged']: notes.append('财务一行=系统两车，系统数为两车相加')
    if s0['id'] == 19: notes.append('王荣胜特批：给客户付按13.8%（×0.862），系统已设本车费率')
    if x['f']['path'] != '带票' and x['f']['src'] == '俊兴' and False: pass
    w5.cell(r, col, '；'.join(notes))
    style_row(w5, r, col, money=tuple(range(9, col)), num=(7,), wrap=(col, 8))
    for c in range(9, col, 3): w5.cell(r, c).fill = C_FIN; w5.cell(r, c + 1).fill = C_SYS
    r += 1
L5 = r - 1; T5 = r
w5.cell(T5, 1, '合计').font = FB
for c in range(9, 9 + 3 * len(STEPS)):
    cc = w5.cell(T5, c, f'=SUM({CL(c)}3:{CL(c)}{L5})'); cc.font = FB; cc.fill = C_GRY; cc.number_format = MONEY
cc = w5.cell(T5, 7, f'=SUM(G3:G{L5})'); cc.font = FB; cc.number_format = '#,##0.00'
for c in range(11, 9 + 3 * len(STEPS), 3):
    rng = f'{CL(c)}3:{CL(c)}{L5}'
    w5.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({CL(c)}3),ABS({CL(c)}3)>1)'], fill=C_RED))
    w5.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({CL(c)}3),ABS({CL(c)}3)>0.005,ABS({CL(c)}3)<=1)'], fill=C_INT, font=Font(name=FNT, size=9, color='98A2B3')))
w5.freeze_panes = 'I3'; w5.auto_filter.ref = f'A2:{CL(8 + 3 * len(STEPS) + 1)}{L5}'
NETDIFF_COL = CL(9 + 3 * (len(STEPS) - 1) + 2)

# ====================== ⑥ 路径统计 ======================
w6 = wb.create_sheet('⑥路径统计')
w6['A1'] = '按计算路径统计（同步后；财务利润按财务表，系统利润按系统；9/21 未录那车只在财务列）'; w6['A1'].font = FB
header(w6, ['路径', '票源', '车数', '吨位', '结算额', '财务利润', '系统利润', '差额(系统−财务)', '每吨利润(系统)', '说明'], row=3,
       widths=[10, 14, 7, 11, 16, 14, 14, 14, 12, 50])
grp = defaultdict(lambda: dict(n=0, ton=0, amt=0, fp=0, sp=0))
for x in R:
    f = x['f']; path = f['path']
    if x['missing']:
        g = grp[(path, '反向-闫东朋' if path == '反向' else '')]; g['fp'] += f['profit']; g['miss'] = 1; continue
    subs = x['subs'] if x['merged'] else [x['s']]
    src = '自带票' if path == '带票' else (subs[0]['src'] or '')
    g = grp[(path, src)]
    g['n'] += len(subs); g['ton'] += sum(s['ton'] for s in subs); g['amt'] += sum(s['amt'] for s in subs)
    g['fp'] += f['profit']; g['sp'] += sum(s['net'] for s in subs)
EXPL = {'配票': '企业票源差价模型：利润 = 量补 − (给客户付 − 前端实付) − 提成', '反向': '反向个人开票：不扣垫资，前端实付=结算额×(1−13.8%)',
        '带票': '供货厂自带票服务费：40 − 垫资23 − 提成15 + 量补25 = 27元/吨'}
r = 4; order = ['配票', '反向', '带票']
for p in order:
    keys = sorted(k for k in grp if k[0] == p)
    first = r
    for k in keys:
        g = grp[k]
        vals = [p, k[1] or '—', g['n'], g['ton'], g['amt'], g['fp'], g['sp'], f'=G{r}-F{r}', f'=IF(D{r}=0,"",G{r}/D{r})',
                '含 9/21 未录一车（财务利润 840.59）' if g.get('miss') else '']
        for j, v in enumerate(vals, 1): w6.cell(r, j, v)
        style_row(w6, r, 10, money=(5, 6, 7, 8, 9), num=(4,), wrap=(10,))
        r += 1
    vals = [f'{p} 小计', '', f'=SUM(C{first}:C{r-1})', f'=SUM(D{first}:D{r-1})', f'=SUM(E{first}:E{r-1})', f'=SUM(F{first}:F{r-1})', f'=SUM(G{first}:G{r-1})',
            f'=G{r}-F{r}', f'=IF(D{r}=0,"",G{r}/D{r})', EXPL[p]]
    for j, v in enumerate(vals, 1): w6.cell(r, j, v)
    style_row(w6, r, 10, money=(5, 6, 7, 8, 9), num=(4,), wrap=(10,))
    for j in range(1, 11): w6.cell(r, j).fill = C_GRY; w6.cell(r, j).font = FB
    r += 1
subtot_rows = [rr for rr in range(4, r) if str(w6.cell(rr, 1).value).endswith('小计')]
vals = ['总计', '', '=' + '+'.join(f'C{x}' for x in subtot_rows), '=' + '+'.join(f'D{x}' for x in subtot_rows), '=' + '+'.join(f'E{x}' for x in subtot_rows),
        '=' + '+'.join(f'F{x}' for x in subtot_rows), '=' + '+'.join(f'G{x}' for x in subtot_rows), f'=G{r}-F{r}', f'=IF(D{r}=0,"",G{r}/D{r})',
        '差额 = 9/21 未录一车(−840.59) + 财务 INT 取整(约 −45)']
for j, v in enumerate(vals, 1): w6.cell(r, j, v)
style_row(w6, r, 10, money=(5, 6, 7, 8, 9), num=(4,), wrap=(10,))
for j in range(1, 11): w6.cell(r, j).font = FB; w6.cell(r, j).fill = C_BLU
TOT6 = r

# 按调货人
r += 3
w6.cell(r - 1, 1, '按调货人统计（同步后，系统）').font = FB
header(w6, ['调货人', '路径', '车数', '吨位', '结算额', '财务利润', '系统利润', '差额', '每吨利润(系统)', '提成(系统)'], row=r)
g2 = defaultdict(lambda: dict(n=0, ton=0, amt=0, fp=0, sp=0, cm=0))
for x in R:
    if x['missing']: continue
    subs = x['subs'] if x['merged'] else [x['s']]
    share = 1 / len(subs)
    for s in subs:
        g = g2[(s['disp'] or '（无）', x['f']['path'])]
        g['n'] += 1; g['ton'] += s['ton']; g['amt'] += s['amt']; g['sp'] += s['net']; g['cm'] += s['comm'] or 0
        g['fp'] += x['f']['profit'] * (s['net'] / sum(t['net'] for t in subs) if x['merged'] else 1)
r += 1; f2 = r
for k in sorted(g2, key=lambda k: (-sum(v['ton'] for kk, v in g2.items() if kk[0] == k[0]), k)):
    g = g2[k]
    vals = [k[0], k[1], g['n'], g['ton'], g['amt'], g['fp'], g['sp'], f'=G{r}-F{r}', f'=IF(D{r}=0,"",G{r}/D{r})', g['cm']]
    for j, v in enumerate(vals, 1): w6.cell(r, j, v)
    style_row(w6, r, 10, money=(5, 6, 7, 8, 9, 10), num=(4,))
    r += 1
w6.cell(r, 1, '合计').font = FB
for c in 'CDEFGHJ':
    cc = w6[f'{c}{r}']; cc.value = f'=SUM({c}{f2}:{c}{r-1})'; cc.font = FB; cc.fill = C_GRY; cc.number_format = MONEY if c not in 'CD' else '#,##0.00'
w6.cell(r + 1, 1, '注：财务 8/29 那行是两车合并，按系统两车利润比例分摊到各自调货人。').font = FG

# ====================== ⑦ 计算过程说明（各路径公式+实例） ======================
w7 = wb.create_sheet('⑦各路径计算过程')
w7.column_dimensions['A'].width = 22; w7.column_dimensions['B'].width = 46; w7.column_dimensions['C'].width = 46; w7.column_dimensions['D'].width = 16; w7.column_dimensions['E'].width = 16
row = 1
EX = {'配票': ('林凯', 50), '反向': ('闫东鹏', 3), '带票': ('林凯-君盛', 3)}
STEPTXT = {
    '配票': [('结算额 A', '应收金额 D = 吨位×单价', '结算金额（磅重明细合计）', 'A'),
             ('垫资费', '23 × 吨位', '吨位 × 户头垫资费率(费率与关系，23)', 'adv'),
             ('前端开票额', '欧冶 E = D − 23×吨位', '结算额 − 垫资费', 'E'),
             ('给客户付', '实际付款 F = INT(D × 0.86)', '结算额 × (1 − 对外点位14%)', 'F'),
             ('前端实付', 'E × 0.865(俊兴/益周) 或 × 0.868(煌烽)', '前端开票额 × (1 − 票源点位)', 'G'),
             ('个人垫资', 'F − 前端实付', '给客户付 − 前端实付', 'P'),
             ('量补', '25 × 吨位（7/26~8/25 那期 20）', '钢厂量补阶梯（那期固化 20）', 'rw'),
             ('调货提成', '林凯 5 × 吨位（9/04 前）', '费率与关系：林凯配票 5 元/吨到 9/04，9/05 起 0', 'cm'),
             ('利润', '前端实付 − F + 量补 − 提成', '量补 − 个人垫资 − 提成 − 其他成本', 'net')],
    '反向': [('结算额 A', '应收金额 D', '结算金额', 'A'), ('垫资费', '不扣（本期已含在反向票点里）', '户头+反向 生效档 = 0', 'adv'),
             ('前端开票额', '就是 D（不扣垫资）', '结算额 − 0', 'E'), ('给客户付', 'F = INT(D × 0.86)', '结算额 × (1 − 14%)', 'F'),
             ('前端实付', '带票公司付 G = D × 0.862', '前端开票额 × (1 − 13.8%)', 'G'), ('个人垫资', 'F − G', '给客户付 − 前端实付', 'P'),
             ('量补', '25 × 吨位', '钢厂量补阶梯', 'rw'), ('调货提成', '0', '0', 'cm'), ('利润', 'G − F + 25×吨位', '量补 − 个人垫资 − 提成', 'net')],
    '带票': [('结算额 A', '应收金额 D', '结算金额', 'A'), ('垫资', '23 × 吨位', '吨位 × 垫资费率 23', 'adv'),
             ('欧冶', 'E = D − 23×吨位', '结算额 − 垫资', 'E'), ('付供货厂', 'F = INT(D − 40×吨位)', '结算额 − 毛服务费', 'F'),
             ('毛服务费', 'D − F ≈ 40 × 吨位', '吨位 × 40', 'G'), ('量补', '25 × 吨位', '钢厂量补阶梯', 'rw'),
             ('调货提成', '15 × 吨位', '费率与关系：林凯带票 15 元/吨', 'cm'), ('利润', 'E − F + 25×吨位 − 15×吨位（=27元/吨）', '毛服务费 − 垫资 + 量补 − 提成', 'net')],
}
for p in ('配票', '反向', '带票'):
    sh, rw_ = EX[p]
    x = next(x for x in R if x['f']['sheet'] == sh and x['f']['row'] == rw_)
    fd = fin_steps(x); sd = sys_steps([x['s']])
    c = w7.cell(row, 1, f"{p}路径　实例：{sh} 第{rw_}行 {x['s']['plate']} {x['f']['date']} {x['f']['ton']}吨（系统 #{x['s']['id']}）"); c.font = Font(name=FNT, bold=True, size=11)
    row += 1
    header(w7, ['步骤', '财务表怎么算', '系统怎么算', '财务结果', '系统结果'], row=row, h=22)
    row += 1
    for lab, ft, st, key in STEPTXT[p]:
        vals = [lab, ft, st, fd.get(key), sd.get(key)]
        for j, v in enumerate(vals, 1):
            cc = w7.cell(row, j, v); cc.font = FB if lab == '利润' else F; cc.border = B; cc.alignment = Alignment(wrap_text=True, vertical='top')
            if j >= 4: cc.number_format = MONEY
        w7.cell(row, 4).fill = C_FIN; w7.cell(row, 5).fill = C_SYS
        row += 1
    row += 2
w7.cell(row, 1, '规则说明（本次统一）').font = FB
for i, t in enumerate(['切割费：由客户承担。钢厂结算单里已经从结算额扣掉，给客户付款也按扣后的结算额算，所以利润里不再扣（以前系统重复扣了，已改）。',
                       '王荣胜 8/29 鲁济宁货4095：特批按 13.8% 给客户付款（×0.862），系统在这一车的“本车费率”里设了。',
                       '林凯配票提成：9/04 以前 5 元/吨，9/05 起停。带票 15 元/吨不变。',
                       '量补：7/26~8/25 那期按约定 20 元/吨（已写进这 10 车的本车费率固定下来），8/26 起按阶梯 25 元/吨。',
                       '财务“实际付款”用 INT() 向下取整，每车差几毛钱，系统不取整；对比表里标灰“INT”，可忽略。'], 1):
    c = w7.cell(row + i, 1, f'{i}. {t}'); c.font = F; w7.merge_cells(start_row=row + i, start_column=1, end_row=row + i, end_column=5)
    c.alignment = Alignment(wrap_text=True, vertical='top'); w7.row_dimensions[row + i].height = 30

# ====================== ⑧ 同步明细 ======================
w8 = wb.create_sheet('⑧系统同步明细')
header(w8, ['车次', '车号', '结算日期', '改了什么（改前 → 改后）', '同步前系统利润', '同步后系统利润', '利润变化', '同步前票差留存', '同步后票差留存', '原因'],
       widths=[7, 14, 11, 60, 14, 14, 12, 13, 13, 40])
r = 2
for tid in sorted(S):
    c = CHG.get(tid); b, a = SB[tid], S[tid]
    txt = '；'.join(f"{k}：{v1} → {v2}" for k, v1, v2 in c['changes']) if c else ''
    reasons = []
    if any(k in ('调货人', '供货厂', '票源', '链条规则') for k, _, _ in (c['changes'] if c else [])): reasons.append('按财务标记同步')
    if a['cut']: reasons.append('切割费客户承担，利润不再扣')
    if tid == 19: reasons.append('王荣胜特批 13.8%')
    if tid == 59: reasons.append('去掉以前的 −16310 冲抵')
    if c and any(k == '本车费率' for k, _, _ in c['changes']) and tid <= 10: reasons.append('7/26~8/25 那期量补 20 固定进本车费率（利润不变）')
    if not txt and abs(a['net'] - b['net']) < 0.005: continue
    vals = [f'#{tid}', a['plate'], a['sd'], txt, b['net'], a['net'], f'=F{r}-E{r}', b['resv'] or 0, a['resv'] or 0, '；'.join(reasons)]
    for j, v in enumerate(vals, 1): w8.cell(r, j, v)
    style_row(w8, r, 10, money=(5, 6, 7, 8, 9), wrap=(4, 10))
    r += 1
L8 = r - 1
w8.cell(r, 1, '合计').font = FB
for cL in 'EFGHI':
    cc = w8[f'{cL}{r}']; cc.value = f'=SUM({cL}2:{cL}{L8})'; cc.font = FB; cc.fill = C_GRY; cc.number_format = MONEY
if L8 >= 2: w8.conditional_formatting.add(f'G2:G{L8}', FormulaRule(formula=['ABS(G2)>0.005'], fill=C_ORG))
w8.freeze_panes = 'D2'; w8.auto_filter.ref = f'A1:J{max(L8, 2)}'
w8.cell(r + 2, 1, '档案：供货厂“𠓾景通”改名“鑫景通”；新增调货人“高栾凤”“王荣胜”（提成 0）。链条规则 5 个模板的“其他成本”节点去掉切割费。同步前已自动备份（data/backups/manual_20260922_063133.db），每项修改在「系统 → 操作记录」里可查。').font = FG
w8.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=10)

# ====================== 总览 ======================
ov = wb.create_sheet('总览', 0)
ov['A1'] = '财务「发货记录表6」× 系统（同步后）对账'; ov['A1'].font = Font(name=FNT, bold=True, size=14)
ov['A2'] = '先按财务标记同步了系统（调货人/供货厂/票源/计算路径），并按你确认的规则改了系统（切割费客户承担只扣一次、王荣胜特批13.8%、林凯配票提成9/05停），再逐项对比。'
ov['A2'].font = FG; ov.merge_cells('A2:F2'); ov.row_dimensions[2].height = 30; ov['A2'].alignment = Alignment(wrap_text=True, vertical='top')
header(ov, ['步骤', '对象', '一致', '需关注', '结论'], row=4, widths=[20, 14, 10, 10, 90], h=22)
steps = [
    ('①车次对账', '财务101行', f"=COUNTIF('①车次对账'!K2:K{L1},\"一致*\")", f"=COUNTA('①车次对账'!K2:K{L1})-C5",
     '99 行一一对上（9 行只是日期差 1~2 天，蓝色）。1 行是财务把两车合成一行（8/29 苏DBE075+冀A224E6）；1 行是 9/21 闫东鹏你还没录'),
    ('②调货人/供货厂', '系统车次', f"=COUNTIF('②调货人供货厂'!J2:J{L2},\"一致\")", f"=COUNTA('②调货人供货厂'!J2:J{L2})-C6",
     '已全部按财务同步。剩下只有名字写法“闫东鹏(财务)/闫东朋(系统)”——你告诉我哪个对，我统一改'),
    ('③票源', '系统车次', f"=COUNTIF('③票源'!H2:H{L3},\"一致\")", f"=COUNTA('③票源'!H2:H{L3})-C7",
     '已全部按财务同步。剩下闫东鹏 9/02 苏BXT103 财务没写票源（系统俊兴，利润两边一致）'),
    ('④计算路径', '系统车次', f"=COUNTIF('④计算路径'!H2:H{L4},\"一致\")", f"=COUNTA('④计算路径'!H2:H{L4})-C8", '全部一致'),
    ('⑤计算过程', '逐车逐步', f"=COUNTIF('⑤计算过程对比'!{NETDIFF_COL}3:{NETDIFF_COL}{L5},\"<=1\")-COUNTIF('⑤计算过程对比'!{NETDIFF_COL}3:{NETDIFF_COL}{L5},\"<-1\")",
     f"=COUNTIF('⑤计算过程对比'!{NETDIFF_COL}3:{NETDIFF_COL}{L5},\">1\")+COUNTIF('⑤计算过程对比'!{NETDIFF_COL}3:{NETDIFF_COL}{L5},\"<-1\")",
     '每一步（结算额→垫资→开票额→给客户付→前端实付→个人垫资→量补→提成→利润）都对上，只剩财务 INT 取整的几毛钱（灰色）'),
]
for i, st in enumerate(steps, 5):
    for j, v in enumerate(st, 1):
        c = ov.cell(i, j, v); c.font = F; c.border = B; c.alignment = Alignment(wrap_text=True, vertical='top')
    ov.row_dimensions[i].height = 32
ov['A11'] = '利润总账'; ov['A11'].font = FB
header(ov, ['项目', '金额', '', '', '说明'], row=12, h=22)
items = [('财务利润合计（101行）', f"=SUM('⑥路径统计'!F{TOT6})", '含 9/21 未录那车 840.59；财务表自己的合计公式漏了这一行'),
         ('系统利润合计（101车）', f"=SUM('⑥路径统计'!G{TOT6})", '同步后。同步前是 186,934.98'),
         ('差额（系统−财务）', '=B14-B13', ''),
         ('其中：9/21 未录一车', -840.5876, '你录入后这项消失'),
         ('其中：财务 INT 取整', '=B15-B16', '99 车每车差不到 1 元，可忽略')]
for i, (a, b, c) in enumerate(items, 13):
    ov.cell(i, 1, a); ov.cell(i, 2, b); ov.cell(i, 5, c)
    for j in (1, 2, 3, 4, 5): ov.cell(i, j).font = FB if i == 15 else F; ov.cell(i, j).border = B
    ov.cell(i, 2).number_format = MONEY
ov['A20'] = '颜色说明'; ov['A20'].font = FB
for k, (fl, txt) in enumerate([(C_RED, '红：不一致（差额>1元、字段不同、系统没录）'), (C_ORG, '橙：本次按财务同步改过的地方 / 合并车次'),
                               (C_BLU, '蓝：只是日期差、名字写法不同'), (C_GRN, '绿：一致'), (C_INT, '灰“INT”：财务 INT 取整的几毛钱，忽略'),
                               (C_FIN, '淡橙：财务表的数'), (C_SYS, '淡蓝：系统的数')], 21):
    ov.cell(k, 1).fill = fl; ov.cell(k, 1).border = B
    ov.cell(k, 2, txt).font = F; ov.merge_cells(start_row=k, start_column=2, end_row=k, end_column=5)
ov['A29'] = '其它表'; ov['A29'].font = FB
for k, t in enumerate(['⑥路径统计：按 配票/反向/带票 × 票源 统计车数、吨位、结算额、两边利润、每吨利润；下面按调货人再统计一遍',
                       '⑦各路径计算过程：三条路径各拿一车真实数据，逐步写出财务怎么算、系统怎么算、两边结果',
                       '⑧系统同步明细：这次系统里每一车改了什么、利润前后变化'], 30):
    c = ov.cell(k, 1, t); c.font = F; ov.merge_cells(start_row=k, start_column=1, end_row=k, end_column=5)
ov.freeze_panes = 'A5'

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(OUT, L1, L2, L3, L4, L5)
