import openpyxl, json, re, datetime as dt, signal
signal.signal(signal.SIGPIPE, signal.SIG_DFL)
"""把财务「发货记录表」逐行和系统车次配对。用法：python3 match_ledger.py 财务台账.xlsx 系统车次.json 输出match.json
配对规则：金额(±1元)或吨位(±0.005)至少一个相同、日期差≤7天，按 金额/吨位/日期差 打分贪心配对。"""
import sys
F=sys.argv[1]
wf=openpyxl.load_workbook(F); wv=openpyxl.load_workbook(F,data_only=True)
fin=[]
for ws in wf.worksheets:
    vs=wv[ws.title]
    for r in range(2,min(ws.max_row,400)+1):
        d=vs.cell(r,1).value; t=vs.cell(r,2).value
        if not isinstance(d,dt.datetime) or not t: continue
        g=lambda c: vs.cell(r,c).value
        gf=lambda c: ws.cell(r,c).value
        src = g(11) if ws.max_column>=11 and isinstance(g(11),str) else None
        jf=str(gf(10)); ef=str(gf(5)); ff=str(gf(6)); gg=str(gf(7)); df=str(gf(4))
        if g(7)=='带票': path='带票'
        elif isinstance(g(5),str) and '反向' in g(5): path='反向'
        else: path='配票'
        # 量补单价、提成
        m=re.findall(r'B\d+\*(\d+)|(\d+)\*B\d+', jf)
        nums=[int(a or b) for a,b in m]
        fin.append(dict(sheet=ws.title,row=r,date=d.date().isoformat(),ton=float(t),price=g(3),amt=float(g(4)),oy=g(5),pay=g(6),front=g(7),left=g(8),paydate=g(9).date().isoformat() if isinstance(g(9),dt.datetime) else None,
            profit=float(g(10)),src=src,path=path,jf=jf,ff=ff,gf=gg,df=df,nums=nums))
S=json.load(open(sys.argv[2]))
# 匹配
pairs=[]
for i,f in enumerate(fin):
    for s in S:
        amt_ok=abs((s['amt'] or 0)-f['amt'])<1
        ton_ok=abs((s['ton'] or 0)-f['ton'])<0.005
        if not (amt_ok or ton_ok): continue
        dd=abs((dt.date.fromisoformat(s['sd'])-dt.date.fromisoformat(f['date'])).days) if s['sd'] else 99
        if dd>7: continue
        score=(0 if amt_ok else 5)+(0 if ton_ok else 5)+dd
        pairs.append((score,i,s['id']))
pairs.sort()
fm,sm={},{}
for sc,i,sid in pairs:
    if i in fm or sid in sm: continue
    fm[i]=(sid,sc); sm[sid]=i
Sid={s['id']:s for s in S}
out=[]
for i,f in enumerate(fin):
    sid,sc=fm.get(i,(None,None))
    out.append(dict(f=f,s=Sid.get(sid),score=sc))
unm_sys=[s for s in S if s['id'] not in sm]
json.dump(dict(rows=out,unmatched_sys=unm_sys),open(sys.argv[3],'w'),ensure_ascii=False,default=str)
print('财务行',len(fin),'匹配',len(fm),'财务利润合计',round(sum(f['profit'] for f in fin),2))
tot_s=0
for o in out:
    f,s=o['f'],o['s']
    flag=''
    if s:
        tot_s+=s['net']
        if abs(s['ton']-f['ton'])>0.005: flag+=' 吨位不同'
        if abs(s['amt']-f['amt'])>1: flag+=' 金额不同'
        if s['sd']!=f['date']: flag+=f' 日期{s["sd"][5:]}'
        d=s['net']-f['profit']
    print(f"{f['sheet']:<6}{f['row']:>3} {f['date'][5:]} {f['ton']:>7} {f['amt']:>11.2f} {f['path']} {f['src'] or '':<3} 财务利润{f['profit']:>10.2f} | " +
          (f"#{s['id']:<3}{s['plate']:<10}{s['disp'] or '-'}:{s['base'] or '-'} {s['src'] or '自带'} {s['tpl']} 系统{s['net']:>10.2f} 差{d:>9.2f}{flag}" if s else '未匹配'))
print('匹配上的系统利润合计',round(tot_s,2))
print('系统有、财务没有:')
for s in unm_sys: print('  #',s['id'],s['sd'],s['plate'],s['ton'],s['amt'],s['disp'],s['base'],s['src'],s['tpl'],round(s['net'],2))
